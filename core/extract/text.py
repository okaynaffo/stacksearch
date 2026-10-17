"""md / txt / docx extraction (PLAN.md §6.3, TODO 1.1)."""

from __future__ import annotations

from pathlib import Path

from core.extract.normalize import is_binary, normalize_text


def extract_text_file(path: Path) -> str:
    """Extract clean text from a .md / .txt file."""
    data = path.read_bytes()
    if is_binary(data):
        raise ValueError(f"binary file: {path}")
    return normalize_text(data)


def extract_docx(path: Path) -> str:
    """Extract clean text from a .docx (paragraphs + tables)."""
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return normalize_text("\n".join(parts))


def extract_by_suffix(path: Path) -> str:
    """Dispatch on suffix. Raises ValueError for unsupported/binary input."""
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown", ".txt", ".text"}:
        return extract_text_file(path)
    if suffix == ".docx":
        return extract_docx(path)
    raise ValueError(f"unsupported suffix for text extraction: {suffix} ({path})")
